import sys
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Cm, Mm, Emu

IMG = sys.argv[1]
LOGO = sys.argv[2]
OUT = sys.argv[3]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(20)
sec.left_margin, sec.right_margin = Mm(30), Mm(15)
sec.footer_distance = Mm(10)
TEXT_W = Cm(16.5)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
pf = st.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = 1.5
pf.first_line_indent = Cm(1.25)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.widow_control = True
# automatic hyphenation
settings = doc.settings.element
ah = OxmlElement('w:autoHyphenation')
ah.set(qn('w:val'), 'true')
settings.append(ah)

# ------------------------------------------------------------------ footer
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.first_line_indent = Cm(0)
r = fp.add_run()
for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if tag:
        e = OxmlElement('w:fldChar')
        e.set(qn('w:fldCharType'), tag)
    else:
        e = OxmlElement('w:instrText')
        e.set(qn('xml:space'), 'preserve')
        e.text = txt
    r._r.append(e)


# ----------------------------------------------------------------- helpers
def para(text='', bold=False, italic=False, size=None, align=None, indent=True, spacing=None,
         before=None, after=None, keep=False):
    p = doc.add_paragraph()
    if text:
        add(p, text, bold, italic, size)
    f = p.paragraph_format
    if align is not None:
        f.alignment = align
    if not indent:
        f.first_line_indent = Cm(0)
    if spacing is not None:
        f.line_spacing = spacing
    if before is not None:
        f.space_before = Pt(before)
    if after is not None:
        f.space_after = Pt(after)
    if keep:
        f.keep_with_next = True
    return p


def add(p, text, bold=False, italic=False, size=None, underline=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if underline:
        r.underline = True
    if size:
        r.font.size = Pt(size)
    return r


def tabs(p, *stops):
    for pos, al in stops:
        p.paragraph_format.tab_stops.add_tab_stop(Cm(pos), al)


L, C, R = WD_TAB_ALIGNMENT.LEFT, WD_TAB_ALIGNMENT.CENTER, WD_TAB_ALIGNMENT.RIGHT

fig_no = [0]


def figure(path, caption, ppi=179, max_h=22.0, max_w=16.5):
    w, h = Image.open(path).size
    wc, hc = w / ppi * 2.54, h / ppi * 2.54
    k = min(1, max_w / wc, max_h / hc)
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0, before=6, keep=True)
    p.add_run().add_picture(path, width=Cm(wc * k))
    fig_no[0] += 1
    para(f'Рисунок {fig_no[0]} – {caption}', align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
         before=6, after=6)
    return fig_no[0]


def heading(text, indent=True):
    return para(text, bold=True, indent=indent, before=6, keep=True)


def formula(name):
    w, h = Image.open(f'{IMG}/{name}.png').size
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=4)
    p.add_run().add_picture(f'{IMG}/{name}.png', width=Cm(w / 300 * 2.54 * 0.85))
    return p


def math(omml_inner, align='center'):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=3)
    xml = (f'<m:oMathPara {nsdecls("m", "w")}><m:oMathParaPr><m:jc m:val="{align}"/></m:oMathParaPr>'
           f'<m:oMath>{omml_inner}</m:oMath></m:oMathPara>')
    p._p.append(parse_xml(xml))
    return p


def mr(t):
    return f'<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t xml:space="preserve">{t}</m:t></m:r>'


def frac(n, d):
    return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'


def nary(ch, sub, sup, e):
    return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/><m:limLoc m:val="undOvr"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{e}</m:e></m:nary>')


# --------------------------------------------------------------- title page
tbl = doc.add_table(rows=1, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
widths = (Cm(3.0), Cm(13.5))
for i, w in enumerate(widths):
    tbl.columns[i].width = w
    tbl.rows[0].cells[i].width = w
c0, c1 = tbl.rows[0].cells
p0 = c0.paragraphs[0]
p0.paragraph_format.first_line_indent = Cm(0)
p0.paragraph_format.line_spacing = 1.0
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
p0.add_run().add_picture(LOGO, width=Cm(2.3))
lines = [('Министерство науки и высшего образования Российской Федерации', True, False),
         ('Калужский филиал федерального', False, False),
         ('государственного бюджетного образовательного', False, False),
         ('учреждения высшего образования', False, False),
         ('«Московский государственный технический университет имени Н.Э. Баумана', True, True),
         ('(национальный исследовательский университет)»', True, True),
         ('(КФ МГТУ им. Н.Э. Баумана)', True, True)]
for i, (t, b, it) in enumerate(lines):
    p = c1.paragraphs[0] if i == 0 else c1.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    add(p, t, b, it, 10)
c1.vertical_alignment = 1

para(indent=False)
p = para(indent=False, before=6)
tabs(p, (16.0, L))
add(p, 'ФАКУЛЬТЕТ _ ', True, size=14)
add(p, 'ИУК «Информатика и управление»', True, True, 14, underline=True)
add(p, '\t', size=14, underline=True)
p = para(indent=False, before=12)
tabs(p, (2.9, L), (16.0, L))
add(p, 'КАФЕДРА ', True, size=14)
add(p, '\tИУК5 «Системы обработки информации»\t', True, True, 14, underline=True)
para(indent=False)
para('О Т Ч Е Т', True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
para('УЧЕБНАЯ ПРАКТИКА', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
para('«Учебно-технологический практикум»', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para('Лабораторная работа № 2', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para('«Реализация основных алгоритмических конструкций»', True, align=WD_ALIGN_PARAGRAPH.CENTER,
     indent=False, after=12)


def person(role, fio):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=6)
    tabs(p, (6.0, L), (9.0, L))
    add(p, role + '\t')
    add(p, '\t', underline=True)
    add(p, f'({fio})', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, after=6)
    tabs(p, (9.5, L), (12.0, L))
    add(p, '\t(подпись)\t(Ф.И.О.)', size=8)


person('Студент гр. ИУК5-13Б', 'Амбарцумян Александр Арташесович')
person('Руководитель', 'Кондратьева Светлана Дмитриевна')


def mark(label, under):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=6)
    tabs(p, (4.2, L), (5.7, L), (7.0, L), (9.5, L))
    add(p, label + '\t')
    add(p, '\t', underline=True)
    add(p, 'баллов\t')
    add(p, '\t', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0)
    tabs(p, (under, L))
    add(p, '\t' + ('(дата)' if under > 7 else '(оценка по пятибалльной шкале)'), size=8)


mark('Оценка руководителя', 7.8)
mark('Оценка защиты', 7.8)
mark('Оценка практики', 5.0)
p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=8)
tabs(p, (8.0, L))
add(p, '\tКомиссия:')
for k in range(3):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=10)
    tabs(p, (8.0, L), (16.4, L))
    add(p, '\t')
    add(p, '\t', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0)
    tabs(p, (8.0, L), (11.0, L))
    add(p, '\t(подпись)\t(Ф.И.О.)', size=8)
p = para('Калуга, 2026', align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=18)
p.add_run().add_break(WD_BREAK.PAGE)

# ---------------------------------------------------------------- content
heading('Цель:', indent=False)
para('Сформировать практические навыки по основам построения базовых алгоритмов и их реализации '
     'на языке C++.')
heading('Задачи:', indent=False)
for t in ('1. Овладеть навыками построения базовых алгоритмов.',
          '2. Овладеть навыками выбора оптимальных алгоритмов программирования.',
          '3. Овладеть навыками описания основных этапов построения алгоритмов.',
          '4. Овладеть навыками реализации основных алгоритмических конструкций: линейной, '
          'разветвляющейся и циклической.',
          '5. Изобразить разработанные алгоритмы в виде блок-схем в соответствии с ГОСТ 19.701-90 '
          'и оформить отчет.'):
    para(t)

heading('Вариант 1', indent=False)


def task_header(n):
    heading(f'№ {n}', indent=False)


def results(t, cases):
    heading('Результат работы программы:')
    for k, cap in enumerate(cases, 1):
        figure(f'{IMG}/out{t}_{k}.png', f'Результат работы программы задачи № {t} ({cap})')


FLOW4_FIG = [0]


def flow_code(t):
    heading('Блок-схема алгоритма:')
    FLOW4_FIG[0] = FLOW4_FIG[0] if t != 4 else fig_no[0] + 1
    figure(f'{IMG}/flow{t}.png', f'Блок-схема алгоритма задачи № {t}', ppi=270)
    heading('Код программы:')
    figure(f'{IMG}/code{t}.png', f'Код программы задачи № {t}')


# --- task 1
task_header(1)
para('Даны x, y. Получить:')
formula('formula1')
heading('Ход решения:')
for t in ('1) Ввести с клавиатуры числа x и y. Переменные объявлены вещественными (тип double), так как '
          'числа могут быть дробными.',
          '2) Вычислить значение выражения по формуле: z = (|x| − |y|) / (1 + |x · y|). Модуль числа '
          'вычисляется функцией fabs() из библиотеки cmath.',
          '3) Вывести полученное значение на экран.'):
    para(t)
para('Знаменатель 1 + |x · y| всегда не меньше единицы, поэтому деления на ноль не возникает при любых '
     'значениях x и y. Алгоритм является линейным.')
flow_code(1)
results(1, ['x = 3, y = 2', 'x = −4, y = 1,5', 'x = 0, y = 5'])
para('Проверка: при x = 3, y = 2 получаем z = (3 − 2) / (1 + 6) = 1 / 7 ≈ 0,142857; при x = −4, y = 1,5 '
     'получаем z = (4 − 1,5) / (1 + 6) = 2,5 / 7 ≈ 0,357143; при x = 0, y = 5 получаем '
     'z = (0 − 5) / 1 = −5. Результаты совпадают с выводом программы.')

# --- task 2
task_header(2)
para('Даны 2 числа. Получить среднее арифметическое кубов этих чисел и среднее геометрическое '
     'модулей этих чисел.')
heading('Ход решения:')
for t in ('1) Ввести с клавиатуры два вещественных числа a и b.',
          '2) Вычислить среднее арифметическое кубов: cubeMean = (a³ + b³) / 2.',
          '3) Вычислить среднее геометрическое модулей: geomMean = √(|a| · |b|). Модули берутся для '
          'того, чтобы подкоренное выражение было неотрицательным; корень вычисляется функцией sqrt().',
          '4) Вывести оба значения на экран.'):
    para(t)
flow_code(2)
results(2, ['a = 2, b = 4', 'a = −3, b = 3'])
para('Проверка: (2³ + 4³) / 2 = (8 + 64) / 2 = 36, √(2 · 4) = √8 ≈ 2,82843; '
     '((−3)³ + 3³) / 2 = (−27 + 27) / 2 = 0, √(3 · 3) = 3.')

# --- task 3
task_header(3)
para('Определить, попадает ли заданная точка внутрь заданной области:')
n_area = figure(f'{IMG}/area.png', 'Заданная область', ppi=150)
heading('Ход решения:')
for t in (f'1) Описать область аналитически. Область на рисунке {n_area} ограничена окружностью радиуса 1 '
          'с центром в начале координат и прямой y = −x. Точка принадлежит области, если она лежит '
          'внутри круга (x² + y² ≤ 1) и не ниже прямой (y ≥ −x). Точки границы считаются '
          'принадлежащими области.',
          '2) Ввести с клавиатуры координаты точки x и y.',
          '3) Проверить составное условие x² + y² ≤ 1 и y ≥ −x. Условия объединяются логической '
          'операцией && (и), так как должны выполняться одновременно.',
          '4) Если условие истинно, вывести сообщение «Точка попадает в область», иначе — '
          '«Точка не попадает в область».'):
    para(t)
flow_code(3)
results(3, ['точка (0,3; 0,4)', 'точка (−0,5; 0,6)', 'точка (0,5; −0,7)', 'точка (0,9; 0,9)'])
para('Проверка: для точки (0,3; 0,4) имеем 0,09 + 0,16 = 0,25 ≤ 1 и 0,4 ≥ −0,3 — точка внутри; '
     'для точки (−0,5; 0,6) имеем 0,25 + 0,36 = 0,61 ≤ 1 и 0,6 ≥ 0,5 — точка внутри; точка (0,5; −0,7) '
     'лежит ниже прямой y = −x (−0,7 < −0,5); точка (0,9; 0,9) лежит вне круга (0,81 + 0,81 = 1,62 > 1).')

# --- task 4
task_header(4)
para('Ввести одномерный массив из n элементов. Вычислить сумму всех отрицательных чисел и сумму всех '
     'положительных чисел.')
heading('Ход решения:')
for t in ('1) Ввести с клавиатуры количество элементов n и выделить под массив динамическую память '
          '(операция new).',
          '2) В цикле for ввести n элементов массива.',
          '3) Присвоить суммам отрицательных (negSum) и положительных (posSum) чисел начальное значение 0.',
          '4) Во втором цикле for перебрать все элементы: если a[i] < 0, прибавить элемент к negSum, '
          'иначе если a[i] > 0 — прибавить к posSum. Нулевые элементы не относятся ни к '
          'положительным, ни к отрицательным и пропускаются.',
          '5) Вывести обе суммы на экран и освободить память (операция delete[]).'):
    para(t)
flow_code(4)
results(4, ['n = 6', 'все элементы отрицательные'])
para('Проверка: для массива 3, −5, 7, 0, −2, 4,5 сумма отрицательных чисел равна −5 + (−2) = −7, сумма '
     'положительных: 3 + 7 + 4,5 = 14,5. Для массива из одних отрицательных чисел сумма положительных '
     'равна 0.')

# --- task 5
task_header(5)
para('Дано натуральное число N. Вычислить:')
formula('formula5')
heading('Ход решения:')
for t in ('1) Ввести с клавиатуры натуральное число N.',
          '2) Присвоить сумме s начальное значение 0, а переменной factI, в которой накапливается '
          'факториал i!, — значение 1.',
          '3) Во внешнем цикле for (i от 1 до N) домножить factI на i, получив i!. Произведению p и '
          'переменной factJ (для накопления j!) присвоить значение 1.',
          '4) Во внутреннем цикле for (j от 1 до i) домножить factJ на j, получив j!, и домножить '
          'произведение p на j! / i!.',
          '5) После завершения внутреннего цикла прибавить произведение p к сумме s.',
          '6) После завершения внешнего цикла вывести s на экран.'):
    para(t)
para('Факториалы не вычисляются заново на каждой итерации, а накапливаются: i! = (i − 1)! · i, '
     'j! = (j − 1)! · j. Для переменных выбран тип double, так как факториалы быстро растут, а '
     'отношение j! / i! является дробным.')
flow_code(5)
results(5, ['N = 1', 'N = 3', 'N = 5'])
para('Проверка: при N = 3 получаем s = 1 + (1! · 2!) / (2!)² + (1! · 2! · 3!) / (3!)³ = '
     '1 + 2 / 4 + 12 / 216 = 1 + 0,5 + 0,055556 ≈ 1,55556, что совпадает с выводом программы.')

# --- control questions
heading('Ответы на контрольные вопросы:', indent=False)
qa = [
    ('1. Опишите, из чего состоит схема программы.',
     'Схема программы состоит из символов процесса, указывающих фактические операции обработки данных '
     '(включая символы, определяющие путь с учетом логических условий), линейных символов, указывающих '
     'поток управления, и специальных символов, облегчающих написание и чтение схемы (соединители, '
     'комментарии).'),
    ('2. Дайте определение понятию «алгоритм».',
     'Алгоритм — это конечная последовательность точно определенных действий (инструкций), выполнение '
     'которых приводит к решению поставленной задачи за конечное число шагов. Алгоритм обладает '
     'свойствами дискретности, определенности, результативности, конечности и массовости.'),
    ('3. Укажите базовые структуры алгоритмов.',
     'Базовыми структурами алгоритмов являются следование (линейная структура), ветвление '
     '(if, if…else, switch…case) и цикл (с предусловием, с постусловием и с параметром).'),
    ('4. Раскройте понятие «условный оператор».',
     'Условный оператор — оператор, который выбирает, какое действие выполнить, в зависимости от '
     'истинности логического условия. В C++ это оператор if (условие) { … } else { … }: если '
     'условие истинно, выполняется первая ветвь, иначе — ветвь else (она может отсутствовать).'),
    ('5. Укажите, для чего используется блок процесса в блок-схемах.',
     'Блок «Процесс» (прямоугольник) обозначает выполнение операции или группы операций, в результате '
     'которых изменяется значение, форма представления или расположение данных, например, '
     'присваивание или вычисление по формуле.'),
    ('6. Укажите, для чего используется блок решения в блок-схемах.',
     'Блок «Решение» (ромб) используется для выбора направления выполнения алгоритма в зависимости '
     'от условия. Он имеет один вход и не менее двух выходов; при двух выходах на них подписывается '
     'результат проверки «да» / «нет».'),
    ('7. Укажите, какие параметры должны быть у блоков в блок-схемах алгоритмов.',
     'Размеры символов определяются параметрами a и b: размер a выбирается из ряда 10, 15, 20 мм '
     '(допускается увеличивать на число, кратное 5 мм), размер b равен 1,5a. Записи внутри символа '
     'должны читаться слева направо и сверху вниз, основное направление потока — сверху вниз и слева '
     'направо, линии в неосновном направлении обозначаются стрелками.'),
    ('8. Приведите примеры задачи и блок-схемы к ней, в которой используется оператор выбора.',
     'Пример задачи: по номеру дня недели (от 1 до 7) вывести его название. В блок-схеме из блока '
     '«Решение», внутри которого записана переменная day, выходит несколько линий, подписанных '
     'значениями 1, 2, …, 7; на каждой линии стоит блок вывода соответствующего названия, а '
     'отдельная линия «иначе» ведет к выводу сообщения об ошибке. Все линии затем сходятся перед '
     'блоком «Конец». В программе это реализуется оператором switch (day) { case 1: … default: … }.'),
    ('9. Приведите примеры задачи и блок-схемы к ней, в которой используется циклический оператор.',
     f'Примером служит задача № 4 данного варианта (рисунок {{n4}}): для перебора элементов массива '
     'используется цикл for, изображенный блоком «Подготовка» (шестиугольник) с записью '
     '«i = 0; i < n; i = i + 1»; тело цикла выполняется для каждого значения i, после чего линия '
     'возвращается в блок цикла, а по окончании цикла управление передается следующему блоку. '
     'Вложенные циклы показаны в задаче № 5.'),
]
for q, a in qa:
    para(q, bold=True)
    para(a.replace('{n4}', str(FLOW4_FIG[0])) if '{n4}' in a else a)

heading('Вывод:', indent=False)
para('В ходе выполнения лабораторной работы были сформированы практические навыки по основам '
     'построения базовых алгоритмов. Были разработаны линейные (задачи № 1 и № 2), разветвляющийся '
     '(задача № 3) и циклические (задачи № 4 и № 5) алгоритмы, построены их блок-схемы в соответствии '
     'с ГОСТ 19.701-90, а сами алгоритмы реализованы на языке C++ с использованием арифметических '
     'операций, математических функций библиотеки cmath, условного оператора if…else, цикла for, '
     'в том числе вложенного, и динамического массива. Работа программ проверена на тестовых данных, '
     'результаты совпадают с ручным расчетом.')

doc.save(OUT)
