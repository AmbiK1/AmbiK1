import sys
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Cm, Mm, Emu

IMG = sys.argv[1]
LOGO = sys.argv[2]
OUT = sys.argv[3]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(20)
sec.left_margin, sec.right_margin = Mm(30), Mm(15)
sec.footer_distance = Mm(10)
TEXT_W = Cm(16.5)

st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
pf = st.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = 1.5
pf.first_line_indent = Cm(1.25)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.widow_control = True
# automatic hyphenation
settings = doc.settings.element
ah = OxmlElement('w:autoHyphenation')
ah.set(qn('w:val'), 'true')
settings.append(ah)

# ------------------------------------------------------------------ footer
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.first_line_indent = Cm(0)
r = fp.add_run()
for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if tag:
        e = OxmlElement('w:fldChar')
        e.set(qn('w:fldCharType'), tag)
    else:
        e = OxmlElement('w:instrText')
        e.set(qn('xml:space'), 'preserve')
        e.text = txt
    r._r.append(e)


# ----------------------------------------------------------------- helpers
def para(text='', bold=False, italic=False, size=None, align=None, indent=True, spacing=None,
         before=None, after=None, keep=False):
    p = doc.add_paragraph()
    if text:
        add(p, text, bold, italic, size)
    f = p.paragraph_format
    if align is not None:
        f.alignment = align
    if not indent:
        f.first_line_indent = Cm(0)
    if spacing is not None:
        f.line_spacing = spacing
    if before is not None:
        f.space_before = Pt(before)
    if after is not None:
        f.space_after = Pt(after)
    if keep:
        f.keep_with_next = True
    return p


def add(p, text, bold=False, italic=False, size=None, underline=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if underline:
        r.underline = True
    if size:
        r.font.size = Pt(size)
    return r


def tabs(p, *stops):
    for pos, al in stops:
        p.paragraph_format.tab_stops.add_tab_stop(Cm(pos), al)


L, C, R = WD_TAB_ALIGNMENT.LEFT, WD_TAB_ALIGNMENT.CENTER, WD_TAB_ALIGNMENT.RIGHT

fig_no = [0]


def figure(path, caption, ppi=179, max_h=22.0, max_w=16.5):
    w, h = Image.open(path).size
    wc, hc = w / ppi * 2.54, h / ppi * 2.54
    k = min(1, max_w / wc, max_h / hc)
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.0, before=6, keep=True)
    p.add_run().add_picture(path, width=Cm(wc * k))
    fig_no[0] += 1
    para(f'Рисунок {fig_no[0]} – {caption}', align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
         before=6, after=6)
    return fig_no[0]


def heading(text, indent=True):
    return para(text, bold=True, indent=indent, before=6, keep=True)


def formula(name):
    w, h = Image.open(f'{IMG}/{name}.png').size
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=4)
    p.add_run().add_picture(f'{IMG}/{name}.png', width=Cm(w / 300 * 2.54 * 0.85))
    return p


def math(omml_inner, align='center'):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=3)
    xml = (f'<m:oMathPara {nsdecls("m", "w")}><m:oMathParaPr><m:jc m:val="{align}"/></m:oMathParaPr>'
           f'<m:oMath>{omml_inner}</m:oMath></m:oMathPara>')
    p._p.append(parse_xml(xml))
    return p


def mr(t):
    return f'<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t xml:space="preserve">{t}</m:t></m:r>'


def frac(n, d):
    return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'


def nary(ch, sub, sup, e):
    return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/><m:limLoc m:val="undOvr"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{e}</m:e></m:nary>')


# --------------------------------------------------------------- title page
tbl = doc.add_table(rows=1, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
widths = (Cm(3.0), Cm(13.5))
for i, w in enumerate(widths):
    tbl.columns[i].width = w
    tbl.rows[0].cells[i].width = w
c0, c1 = tbl.rows[0].cells
p0 = c0.paragraphs[0]
p0.paragraph_format.first_line_indent = Cm(0)
p0.paragraph_format.line_spacing = 1.0
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
p0.add_run().add_picture(LOGO, width=Cm(2.3))
lines = [('Министерство науки и высшего образования Российской Федерации', True, False),
         ('Калужский филиал федерального', False, False),
         ('государственного бюджетного образовательного', False, False),
         ('учреждения высшего образования', False, False),
         ('«Московский государственный технический университет имени Н.Э. Баумана', True, True),
         ('(национальный исследовательский университет)»', True, True),
         ('(КФ МГТУ им. Н.Э. Баумана)', True, True)]
for i, (t, b, it) in enumerate(lines):
    p = c1.paragraphs[0] if i == 0 else c1.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    add(p, t, b, it, 10)
c1.vertical_alignment = 1

para(indent=False)
p = para(indent=False, before=6)
tabs(p, (16.0, L))
add(p, 'ФАКУЛЬТЕТ _ ', True, size=14)
add(p, 'ИУК «Информатика и управление»', True, True, 14, underline=True)
add(p, '\t', size=14, underline=True)
p = para(indent=False, before=12)
tabs(p, (2.9, L), (16.0, L))
add(p, 'КАФЕДРА ', True, size=14)
add(p, '\tИУК5 «Системы обработки информации»\t', True, True, 14, underline=True)
para(indent=False)
para('О Т Ч Е Т', True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
para('УЧЕБНАЯ ПРАКТИКА', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
para('«Учебно-технологический практикум»', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para('Лабораторная работа № 3', True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para('«Действия над массивами и строками»', True, align=WD_ALIGN_PARAGRAPH.CENTER,
     indent=False, after=12)


def person(role, fio):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=6)
    tabs(p, (6.0, L), (9.0, L))
    add(p, role + '\t')
    add(p, '\t', underline=True)
    add(p, f'({fio})', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, after=6)
    tabs(p, (9.5, L), (12.0, L))
    add(p, '\t(подпись)\t(Ф.И.О.)', size=8)


person('Студент гр. ИУК5-13Б', 'Амбарцумян Александр Арташесович')
person('Руководитель', 'Кондратьева Светлана Дмитриевна')


def mark(label, under):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=6)
    tabs(p, (4.2, L), (5.7, L), (7.0, L), (9.5, L))
    add(p, label + '\t')
    add(p, '\t', underline=True)
    add(p, 'баллов\t')
    add(p, '\t', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0)
    tabs(p, (under, L))
    add(p, '\t' + ('(дата)' if under > 7 else '(оценка по пятибалльной шкале)'), size=8)


mark('Оценка руководителя', 7.8)
mark('Оценка защиты', 7.8)
mark('Оценка практики', 5.0)
p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=8)
tabs(p, (8.0, L))
add(p, '\tКомиссия:')
for k in range(3):
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=10)
    tabs(p, (8.0, L), (16.4, L))
    add(p, '\t')
    add(p, '\t', underline=True)
    p = para(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0)
    tabs(p, (8.0, L), (11.0, L))
    add(p, '\t(подпись)\t(Ф.И.О.)', size=8)
p = para('Калуга, 2026', align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=18)
p.add_run().add_break(WD_BREAK.PAGE)

# ---------------------------------------------------------------- content
heading('Цель:', indent=False)
para('Сформировать практические навыки разработки эффективных алгоритмов обработки массивов и строк '
     'и их реализации на языке C++.')
heading('Задачи:', indent=False)
for t in ('1. Овладеть навыками построения эффективных алгоритмов обработки массивов и строк.',
          '2. Овладеть навыками выбора оптимальных алгоритмов программирования.',
          '3. Овладеть навыками описания основных этапов построения алгоритмов.',
          '4. Изобразить разработанные алгоритмы в виде блок-схем в соответствии с ГОСТ 19.701-90.',
          '5. Реализовать алгоритмы на языке C++ и оформить отчет.'):
    para(t)

heading('Вариант 1', indent=False)


def task_header(n):
    heading(f'№ {n}', indent=False)


def results(t, cases):
    heading('Результат работы программы:')
    for k, cap in enumerate(cases, 1):
        figure(f'{IMG}/out{t}_{k}.png', f'Результат работы программы задачи № {t} ({cap})')


FIG = {}


def flow_code(t):
    heading('Блок-схема алгоритма:')
    FIG[t] = fig_no[0] + 1
    figure(f'{IMG}/flow{t}.png', f'Блок-схема алгоритма задачи № {t}', ppi=270, max_h=23)
    heading('Код программы:')
    figure(f'{IMG}/code{t}.png', f'Код программы задачи № {t}')


def steps(*items):
    for t in items:
        para(t)


# --- task 1
task_header(1)
para('Дана матрица A(n × m). Найти порядковые номера первого отрицательного и последнего '
     'положительного элемента (если таковые имеются). Значение элементов и их порядковые номера '
     'вывести на экран или выдать соответствующее сообщение.')
heading('Ход решения:')
steps('1) Ввести размеры матрицы n и m, а затем ее элементы. Матрица объявлена как двумерный массив '
      'double a[MAX][MAX], где MAX = 10 — максимальный размер.',
      '2) Порядковым номером элемента считается его номер при просмотре матрицы по строкам слева '
      'направо, начиная с 1. Для элемента a[i][j] он равен i · m + j + 1.',
      '3) Переменным firstNeg и lastPos присвоить значение −1 — признак того, что нужный элемент '
      'еще не найден.',
      '4) Двумя вложенными циклами for просмотреть все элементы. Если элемент отрицательный и '
      'firstNeg еще равен −1, запомнить его номер в firstNeg (так запоминается только первый '
      'отрицательный элемент). Если элемент положительный, запомнить его номер в lastPos (номер '
      'перезаписывается, поэтому в конце останется последний положительный элемент).',
      '5) Если firstNeg ≠ −1, вывести элемент и его порядковый номер, иначе — сообщение об '
      'отсутствии отрицательных элементов. Аналогично для lastPos. Строка и столбец элемента '
      'восстанавливаются по номеру: строка = номер / m, столбец = номер mod m.')
para('Из-за размера блок-схема разделена на две части, связанные соединителем «A».')
flow_code(1)
results(1, ['есть оба элемента', 'нет отрицательных', 'нет положительных'])
para('Проверка: в первой матрице по строкам идут элементы 2, 5, −1, 0, −3, 7, 4, −6, 1, −2, 8, 0. '
     'Первый отрицательный — −1 с номером 3, последний положительный — 8 с номером 11. '
     'Во втором и третьем примерах программа выдает соответствующие сообщения.')

# --- task 2
task_header(2)
para('Дан массив A(n). Построить матрицу A(n × n) вида:')
p = para(indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=4)
p.add_run().add_picture(f'{IMG}/matrix2.png', width=Cm(5.5))
heading('Ход решения:')
steps('1) Ввести n и элементы массива a.',
      '2) Заметим, что каждая следующая строка матрицы — это предыдущая строка, циклически '
      'сдвинутая на один элемент влево. Поэтому элемент, стоящий в строке i и столбце j (нумерация '
      'с 0), равен a[(i + j) mod n]: остаток от деления на n «заворачивает» индекс в начало массива.',
      '3) Двумя вложенными циклами заполнить матрицу b[i][j] = a[(i + j) mod n].',
      '4) Вывести матрицу построчно; для выравнивания столбцов используется манипулятор setw(5) '
      'из библиотеки iomanip.')
flow_code(2)
results(2, ['n = 5', 'n = 3'])
para('Проверка: при n = 5 вторая строка равна 2, 3, 4, 5, 1, последняя — 5, 1, 2, 3, 4, что '
     'соответствует заданному виду матрицы.')

# --- task 3
task_header(3)
para('Характеристикой столбца целочисленной матрицы A(n × m) назовем сумму модулей его отрицательных '
     'нечетных элементов. Переставляя столбцы заданной матрицы, расположить их в соответствии с '
     'ростом их характеристик.')
heading('Ход решения:')
steps('1) Ввести n, m и элементы целочисленной матрицы.',
      '2) Для каждого столбца j вычислить характеристику ch[j]: обнулить ее и просмотреть элементы '
      'столбца; если элемент отрицательный и нечетный (a[i][j] % 2 ≠ 0), прибавить его модуль '
      '(функция abs()).',
      '3) Упорядочить столбцы методом пузырька по массиву характеристик: на каждом проходе p '
      'сравниваются соседние характеристики ch[j] и ch[j + 1], и если ch[j] > ch[j + 1], то '
      'меняются местами сами характеристики и все элементы столбцов j и j + 1 (функция std::swap). '
      'После каждого прохода самый «тяжелый» столбец оказывается в конце, поэтому внутренний цикл '
      'идет до m − 1 − p.',
      '4) Вывести упорядоченные характеристики и полученную матрицу.')
para('Блок-схема разделена на две части, связанные соединителем «A».')
flow_code(3)
results(3, ['n = 3, m = 4'])
para('Проверка: характеристики исходных столбцов: |−5| = 5; |−9| + |−3| = 12; |−3| + |−7| = 10; '
     '|−1| = 1 (элемент −2 четный и не учитывается). По возрастанию характеристик столбцы идут в '
     'порядке 4, 1, 3, 2, что совпадает с результатом программы.')

# --- task 4
task_header(4)
para('Из текста удалить все слова заданной длины, начинающиеся с согласных букв.')
heading('Ход решения:')
steps('1) Ввести текст в символьный массив text функцией cin.getline() (она, в отличие от cin >>, '
      'читает строку вместе с пробелами) и длину слова k. Текст вводится латинскими буквами; '
      'гласными считаются буквы a, e, i, o, u, y.',
      '2) Результат формируется в отдельном массиве result; i — индекс в исходном тексте, '
      'r — в результате.',
      '3) Пока не достигнут конец строки (символ \'\\0\'), проверять текущий символ. Если это не '
      'буква (пробел, знак препинания), скопировать его в result.',
      '4) Если это буква — начало слова: запомнить start = i и сдвигать i, пока идут буквы. Длина '
      'слова len = i − start. Первая буква согласная, если ее нет в строке гласных (функция strchr() '
      'возвращает NULL).',
      '5) Если len = k и слово начинается с согласной, слово не копируется, а следующие за ним '
      'пробелы пропускаются, чтобы в тексте не оставалось двойных пробелов. Иначе слово '
      'посимвольно копируется в result.',
      '6) В конце дописать в result нулевой символ и вывести результат.')
flow_code(4)
results(4, ['k = 3', 'k = 3, несколько удаляемых слов'])
para('Проверка: в первом тексте слова длины 3 — big, red, and; удалены big и red, а слово and '
     'оставлено, так как начинается с гласной. Во втором тексте удалены The, cat, the, dog, а '
     'and и old оставлены.')

# --- task 5
task_header(5)
para('Найти, каких букв в тексте больше — гласных или согласных.')
heading('Ход решения:')
steps('1) Ввести текст функцией cin.getline(). Текст вводится латинскими буквами; гласными считаются '
      'a, e, i, o, u, y в любом регистре.',
      '2) Обнулить счетчики vowels (гласные) и consonants (согласные).',
      '3) В цикле for перебрать символы до нулевого символа. Если символ — буква (функция isalpha()), '
      'проверить, есть ли он в строке гласных (функция strchr()): если есть, увеличить vowels, '
      'иначе — consonants. Пробелы, цифры и знаки препинания не учитываются.',
      '4) Вывести оба количества и сравнить их: вывести, каких букв больше, или сообщение о том, '
      'что их поровну.')
flow_code(5)
results(5, ['согласных больше', 'гласных больше', 'поровну'])
para('Проверка: в тексте «Data base» гласные a, a, a, e (4), согласные D, t, b, s (4) — поровну.')

# --- control questions
heading('Ответы на контрольные вопросы:', indent=False)
qa = [
    ('1. Дайте определение понятию «массив».',
     'Массив — это совокупность определенного количества однотипных переменных, имеющих одно имя и '
     'расположенных в памяти последовательно. Например, int a[3] — массив из трех переменных '
     'типа int.'),
    ('2. Укажите, как называют переменные массива.',
     'Переменные массива называют элементами массива.'),
    ('3. Укажите, с помощью чего можно получить доступ к элементу массива.',
     'Доступ к элементу массива выполняется по индексу — порядковому номеру элемента, который '
     'указывается в квадратных скобках после имени массива, например a[0] = 33. Индексация '
     'начинается с 0, поэтому последний элемент массива из n элементов имеет индекс n − 1. Для '
     'двумерного массива указываются два индекса: a[i][j].'),
    ('4. Приведите примеры объявления массивов.',
     'int a[10]; — массив из 10 целых чисел без инициализации; int a[5] = {1, 2, 3, 4, 5}; — с '
     'инициализацией; int a[] = {1, 2, 3}; — размер вычисляет компилятор; int a[10] = {}; — все '
     'элементы равны 0; int a[10] {1, 2}; — первые два элемента 1 и 2, остальные 0; '
     'double m[3][4]; — двумерный массив (матрица) из 3 строк и 4 столбцов.'),
    ('5. Раскройте значение термина «строка».',
     'Строка в C++ — это массив символов типа char, последним элементом которого является нулевой '
     'символ \'\\0\'. Именно нулевой символ отмечает конец строки и позволяет работать с массивом '
     'как с единым текстом.'),
    ('6. Укажите различия массивов и строк.',
     'Строка — частный случай массива: она хранит символы (тип char) и обязательно заканчивается '
     'символом \'\\0\'. Поэтому размер массива под строку на 1 больше числа символов. Строку можно '
     'вывести целиком по имени (cout << st), ввести функциями cin >> и cin.getline() и обрабатывать '
     'функциями библиотеки cstring. Обычный массив выводится и вводится только поэлементно, в '
     'цикле.'),
    ('7. Приведите примеры объявления строк.',
     'char st[] = "hello!"; — строка инициализируется строковой константой, \'\\0\' добавляется '
     'автоматически (размер массива 7); char st[] {\'h\', \'i\', \'\\0\'}; — посимвольная '
     'инициализация с явным нулевым символом; char st[20] = ""; — пустая строка, все элементы '
     'равны \'\\0\'.'),
    ('8. Перечислите и опишите основные операции работы со строками.',
     'Ввод строки: cin >> st (читает до первого пробела) и cin.getline(st, размер) (читает строку '
     'целиком вместе с пробелами); вывод: cout << st. Функции библиотеки cstring: strlen() — длина '
     'строки без \'\\0\'; strcpy() — копирование одной строки в другую; strcat() — объединение '
     '(конкатенация) строк; strcmp() — сравнение строк (возвращает 0, если строки равны); '
     'strchr() — поиск символа в строке. Кроме того, к отдельным символам можно обращаться по '
     'индексу, как к элементам массива.'),
    ('9. Укажите, что необходимо сделать, чтобы присвоить элементу массива значение \'\\0\'.',
     'Нужно обратиться к элементу по индексу и присвоить ему нулевой символ: st[i] = \'\\0\'; '
     '(или st[i] = 0;). Чтобы сразу заполнить нулевыми символами все элементы, массив '
     'инициализируют пустой строкой: char st[20] = "";. Так в задаче № 4 после формирования '
     'результата выполняется result[r] = \'\\0\', чтобы обозначить конец строки.'),
]
for q, a in qa:
    para(q, bold=True, keep=True)
    para(a)

heading('Вывод:', indent=False)
para('В ходе выполнения лабораторной работы были сформированы практические навыки разработки '
     'алгоритмов обработки массивов и строк. Были решены задачи на поиск элементов матрицы по '
     'условию, построение матрицы из одномерного массива с использованием циклического сдвига '
     'индексов, перестановку столбцов матрицы по возрастанию их характеристик методом пузырька, '
     'удаление слов из текста и подсчет гласных и согласных букв. Для работы с текстом использованы '
     'символьные массивы, функция cin.getline() и функции библиотек cstring и cctype. Для всех '
     'алгоритмов построены блок-схемы в соответствии с ГОСТ 19.701-90, работа программ проверена на '
     'тестовых данных, результаты совпадают с ручным расчетом.')

doc.save(OUT)
